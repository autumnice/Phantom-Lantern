"""Export service for generating PPTX and PDF files."""

import base64
import io
import uuid
from pathlib import Path
from typing import List, Optional

from fastapi import HTTPException

from app.config import get_settings
from app.schemas.requests import AspectRatio, ExportSlide
from app.services.errors import create_error_response


class ExportService:
    """Service for exporting presentations to PPTX and PDF."""

    def __init__(self, export_dir: str):
        self.export_dir = Path(export_dir)
        self.export_dir.mkdir(parents=True, exist_ok=True)

    def _get_slide_dimensions(self, aspect_ratio: AspectRatio) -> tuple[float, float]:
        """Get slide dimensions in inches based on aspect ratio."""
        dimensions = {
            "16:9": (13.333, 7.5),
            "4:3": (10.0, 7.5),
            "1:1": (7.5, 7.5),
            "3:4": (7.5, 10.0),
            "9:16": (7.5, 13.333),
        }
        return dimensions.get(aspect_ratio, (13.333, 7.5))

    def _decode_image_data(self, image_data: str) -> bytes:
        """Decode base64 image data."""
        if image_data.startswith("data:"):
            base64_part = image_data.split(",", 1)[1] if "," in image_data else image_data
        else:
            base64_part = image_data
        return base64.b64decode(base64_part)

    def _generate_filename(self, extension: str) -> tuple[str, Path]:
        """Generate unique filename and filepath."""
        filename = f"presentation_{uuid.uuid4().hex[:8]}.{extension}"
        return filename, self.export_dir / filename

    def _get_image_stream(self, slide_dict: dict) -> Optional[io.BytesIO]:
        """Extract image stream from slide data, returns None if no image or error."""
        if not slide_dict.get("imageBase64"):
            return None
        try:
            image_bytes = self._decode_image_data(slide_dict["imageBase64"])
            return io.BytesIO(image_bytes)
        except Exception as e:
            print(f"Error decoding image: {e}")
            return None

    async def export_pptx(
        self,
        slides: List[ExportSlide],
        aspect_ratio: AspectRatio,
        request_id: Optional[str] = None,
    ) -> str:
        """Export slides to PPTX format."""
        try:
            from pptx import Presentation
            from pptx.util import Inches, Pt
            from pptx.enum.text import PP_ALIGN

            width, height = self._get_slide_dimensions(aspect_ratio)
            prs = Presentation()
            prs.slide_width = Inches(width)
            prs.slide_height = Inches(height)
            blank_layout = prs.slide_layouts[6]

            for slide_data in slides:
                slide_dict = slide_data.model_dump()
                slide = prs.slides.add_slide(blank_layout)

                image_stream = self._get_image_stream(slide_dict)
                if image_stream:
                    slide.shapes.add_picture(image_stream, Inches(0), Inches(0), width=Inches(width), height=Inches(height))

                title = slide_dict.get("title", "")
                if title:
                    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(width - 1), Inches(1))
                    title_frame = title_box.text_frame
                    title_frame.paragraphs[0].text = title
                    title_frame.paragraphs[0].font.size = Pt(36)
                    title_frame.paragraphs[0].font.bold = True
                    title_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

                content = slide_dict.get("content", "")
                if content:
                    content_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.8), Inches(width - 1), Inches(height - 2.5))
                    content_frame = content_box.text_frame
                    content_frame.word_wrap = True
                    content_frame.paragraphs[0].text = content
                    content_frame.paragraphs[0].font.size = Pt(18)

            filename, filepath = self._generate_filename("pptx")
            prs.save(str(filepath))
            return f"/exports/{filename}"

        except Exception as e:
            import traceback
            traceback.print_exc()
            create_error_response("EXPORT_FAILED", f"Failed to export PPTX: {str(e)}", 500, request_id)

    async def export_pdf(
        self,
        slides: List[ExportSlide],
        aspect_ratio: AspectRatio,
        request_id: Optional[str] = None,
    ) -> str:
        """Export slides to PDF format."""
        try:
            from reportlab.lib.units import inch
            from reportlab.pdfgen import canvas
            from reportlab.lib.utils import ImageReader

            width_inches, height_inches = self._get_slide_dimensions(aspect_ratio)
            page_width, page_height = width_inches * inch, height_inches * inch
            filename, filepath = self._generate_filename("pdf")
            c = canvas.Canvas(str(filepath), pagesize=(page_width, page_height))

            for slide_data in slides:
                slide_dict = slide_data.model_dump()

                image_stream = self._get_image_stream(slide_dict)
                if image_stream:
                    c.drawImage(ImageReader(image_stream), 0, 0, width=page_width, height=page_height)

                title = slide_dict.get("title", "")
                if title:
                    c.setFont("Helvetica-Bold", 36)
                    c.drawCentredString(page_width / 2, page_height - 0.75 * inch, title)

                content = slide_dict.get("content", "")
                if content:
                    c.setFont("Helvetica", 14)
                    text_object = c.beginText(0.5 * inch, page_height - 1.8 * inch)
                    for line in content.split("\n"):
                        text_object.textLine(line)
                    c.drawText(text_object)

                c.showPage()

            c.save()
            return f"/exports/{filename}"

        except Exception as e:
            import traceback
            traceback.print_exc()
            create_error_response("EXPORT_FAILED", f"Failed to export PDF: {str(e)}", 500, request_id)


_export_service: Optional[ExportService] = None


def get_export_service() -> ExportService:
    """Get or create the Export service singleton."""
    global _export_service
    if _export_service is None:
        settings = get_settings()
        _export_service = ExportService(export_dir=settings.export_dir)
    return _export_service
